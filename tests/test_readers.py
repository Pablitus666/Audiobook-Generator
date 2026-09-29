
from pathlib import Path

import pytest

from audiobook_generator.core.errors import (
    DocumentReadError,
    InputFileError,
)
from audiobook_generator.readers.docx import DocxReader
from audiobook_generator.readers.factory import ReaderFactory
from audiobook_generator.readers.html import HtmlReader
from audiobook_generator.readers.odt import OdtReader
from audiobook_generator.readers.rtf import RtfReader
from audiobook_generator.readers.pdf import PdfReader
from audiobook_generator.readers.text import TextReader


def test_text_reader_preserves_utf8(tmp_path: Path):
    source = tmp_path / "libro.txt"
    source.write_text(
        "Capítulo 1\n\nÁrbol, corazón y acción.",
        encoding="utf-8",
    )

    document = TextReader().read(source)

    assert document.title == "libro"
    assert document.text == "Capítulo 1\n\nÁrbol, corazón y acción."
    assert document.chapters == []


def test_empty_text_has_no_chapters(tmp_path: Path):
    source = tmp_path / "vacio.txt"
    source.write_text("", encoding="utf-8")

    document = TextReader().read(source)

    assert document.text == ""
    assert document.chapters == []


def test_text_reader_rejects_missing_file(tmp_path: Path):
    source = tmp_path / "inexistente.txt"

    with pytest.raises(
        InputFileError,
        match="no existe el archivo de entrada",
    ):
        TextReader().read(source)


def test_text_reader_rejects_invalid_utf8(tmp_path: Path):
    source = tmp_path / "invalido.txt"

    source.write_bytes(
        b"\xff\xfe\xfa\xfb"
    )

    with pytest.raises(
        DocumentReadError,
        match="no se pudo decodificar el archivo como UTF-8",
    ):
        TextReader().read(source)


def test_pdf_reader_reads_valid_pdf(tmp_path: Path):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    source = tmp_path / "libro.pdf"

    pdf = canvas.Canvas(
        str(source),
        pagesize=A4,
    )

    pdf.drawString(
        72,
        780,
        "Capítulo 1",
    )

    pdf.drawString(
        72,
        750,
        "Texto de prueba para PDF.",
    )

    pdf.save()

    document = PdfReader().read(source)

    assert document.title == "libro"
    assert "Capítulo 1" in document.text
    assert "Texto de prueba para PDF." in document.text


def test_pdf_reader_rejects_missing_file(tmp_path: Path):
    source = tmp_path / "inexistente.pdf"

    with pytest.raises(
        InputFileError,
        match="no existe el archivo de entrada",
    ):
        PdfReader().read(source)


def test_pdf_reader_rejects_corrupt_file(tmp_path: Path):
    source = tmp_path / "corrupto.pdf"

    source.write_bytes(
        b"esto no es un PDF valido"
    )

    with pytest.raises(
        DocumentReadError,
        match="no se pudo leer el PDF",
    ):
        PdfReader().read(source)


def test_docx_reader_reads_valid_docx(tmp_path: Path):
    from docx import Document as DocxDocument

    source = tmp_path / "libro.docx"

    document = DocxDocument()

    document.add_heading(
        "Capítulo 1",
        level=1,
    )

    document.add_paragraph(
        "Texto de prueba para DOCX."
    )

    document.save(source)

    result = DocxReader().read(source)

    assert result.title == "libro"
    assert "Capítulo 1" in result.text
    assert "Texto de prueba para DOCX." in result.text


def test_docx_reader_rejects_missing_file(tmp_path: Path):
    source = tmp_path / "inexistente.docx"

    with pytest.raises(
        InputFileError,
        match="no existe el archivo de entrada",
    ):
        DocxReader().read(source)


def test_docx_reader_rejects_corrupt_file(tmp_path: Path):
    source = tmp_path / "corrupto.docx"

    source.write_bytes(
        b"esto no es un DOCX valido"
    )

    with pytest.raises(
        DocumentReadError,
        match="no se pudo leer el documento DOCX",
    ):
        DocxReader().read(source)


def test_unsupported_format_raises():
    with pytest.raises(
        ValueError,
        match="Formato no soportado",
    ):
        ReaderFactory.create(
            Path("libro.xyz")
        )



def test_rtf_reader_reads_text(tmp_path: Path):
    source = tmp_path / "libro.rtf"
    source.write_bytes(
        br"{\rtf1\ansi{\fonttbl{\f0 Arial;}}"
        br"\f0\fs24 Cap{\'edtulo 1}\par "
        br"Texto con {\b formato} y \u233?}."
    )

    document = RtfReader().read(source)

    assert document.title == "libro"
    assert "Capítulo 1" in document.text
    assert "Texto con formato y" in document.text
    assert "Arial" not in document.text


def test_rtf_reader_rejects_invalid_file(tmp_path: Path):
    source = tmp_path / "corrupto.rtf"
    source.write_bytes(b"no es rtf")

    with pytest.raises(
        DocumentReadError,
        match="no contiene un documento RTF válido",
    ):
        RtfReader().read(source)


def test_html_reader_reads_title_and_text(tmp_path: Path):
    source = tmp_path / "libro.html"
    source.write_text(
        """<!doctype html>
<html>
<head><title>Mi libro</title><style>oculto</style></head>
<body>
<h1>Capítulo 1</h1>
<p>Primer párrafo.</p>
<p>Segundo <strong>párrafo</strong>.</p>
<script>no debe aparecer</script>
</body>
</html>""",
        encoding="utf-8",
    )

    document = HtmlReader().read(source)

    assert document.title == "Mi libro"
    assert "Capítulo 1" in document.text
    assert "Primer párrafo." in document.text
    assert "Segundo párrafo." in document.text
    assert "oculto" not in document.text
    assert "no debe aparecer" not in document.text


def test_html_reader_rejects_missing_file(tmp_path: Path):
    source = tmp_path / "inexistente.html"

    with pytest.raises(
        InputFileError,
        match="no existe el archivo de entrada",
    ):
        HtmlReader().read(source)


def test_odt_reader_reads_text(tmp_path: Path):
    from zipfile import ZIP_DEFLATED, ZipFile

    source = tmp_path / "libro.odt"

    content = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content
 xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
 xmlns:dc="http://purl.org/dc/elements/1.1/">
 <office:body>
  <office:text>
   <text:h text:outline-level="1">Capítulo 1</text:h>
   <text:p>Primer párrafo.</text:p>
   <text:p>Segundo <text:span>párrafo</text:span>.</text:p>
  </office:text>
 </office:body>
</office:document-content>
"""

    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        archive.writestr("content.xml", content)

    document = OdtReader().read(source)

    assert document.title == "libro"
    assert "Capítulo 1" in document.text
    assert "Primer párrafo." in document.text
    assert "Segundo párrafo." in document.text


def test_odt_reader_rejects_corrupt_file(tmp_path: Path):
    source = tmp_path / "corrupto.odt"
    source.write_bytes(b"no es odt")

    with pytest.raises(
        DocumentReadError,
        match="no se pudo leer el documento ODT",
    ):
        OdtReader().read(source)


def test_html_reader_honors_meta_charset(tmp_path: Path):
    source = tmp_path / "libro.html"
    source.write_bytes(
        b"<html><head><meta charset=\"windows-1252\"><title>Libro</title></head>"
        b"<body><p>Cap\xedtulo con coraz\xf3n y acci\xf3n.</p></body></html>"
    )

    document = HtmlReader().read(source)

    assert document.title == "Libro"
    assert "Capítulo con corazón y acción." in document.text


def test_odt_reader_rejects_missing_file(tmp_path: Path):
    source = tmp_path / "inexistente.odt"

    with pytest.raises(
        InputFileError,
        match="no existe el archivo de entrada",
    ):
        OdtReader().read(source)


def test_rtf_reader_rejects_missing_file(tmp_path: Path):
    source = tmp_path / "inexistente.rtf"

    with pytest.raises(
        InputFileError,
        match="no existe el archivo de entrada",
    ):
        RtfReader().read(source)
