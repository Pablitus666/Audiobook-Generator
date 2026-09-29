from __future__ import annotations

from pathlib import Path

from audiobook_generator.core.errors import (
    DocumentReadError,
    InputFileError,
)
from audiobook_generator.core.models import Document
from audiobook_generator.readers.base import DocumentReader


class DocxReader(DocumentReader):
    """Lector de documentos DOCX."""

    def read(self, source: Path) -> Document:
        if not source.is_file():
            raise InputFileError(
                f"no existe el archivo de entrada: {source}"
            )

        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise DocumentReadError(
                "Para DOCX instala las dependencias de documentos: "
                "pip install -r requirements-docs.txt"
            ) from exc

        try:
            document = DocxDocument(str(source))

            paragraphs = [
                paragraph.text.strip()
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

        except Exception as exc:
            raise DocumentReadError(
                f"no se pudo leer el documento DOCX: {source}"
            ) from exc

        return Document(
            title=source.stem,
            source=source,
            text="\n\n".join(paragraphs),
        )